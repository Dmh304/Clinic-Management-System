import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

test_cases = {
    "UC-27d": [
        ["TC_UC27d_01", "Cancel session successfully by Nurse/Staff", "1. Login as Nurse/Staff\n2. View Care Queue\n3. Select a pending session\n4. Click 'Cancel'\n5. Confirm cancellation", "Session is marked as CANCELLED\nSuccess message 'Huỷ buổi khám thành công'", "Session must be in pending/booked state"],
        ["TC_UC27d_02", "Cancel session successfully by Manager", "1. Login as Manager\n2. Navigate to All Care Sessions\n3. Select a session\n4. Click 'Cancel'\n5. Confirm cancellation", "Session is marked as CANCELLED\nSuccess message 'Huỷ buổi khám thành công'", "Session must exist"],
    ],
    "UC-38": [
        ["TC_UC38_01", "View Fabrication Queue", "1. Login as Optician/Staff\n2. Navigate to Eyeglass Fabrication Queue", "List of pending eyeglass orders is displayed\nAPI returns 'Lấy hàng đợi gia công kính thành công'", "Orders must exist in fabrication queue"],
        ["TC_UC38_02", "Start Fabrication process", "1. Login as Optician\n2. Go to Fabrication Queue\n3. Select an order\n4. Click 'Start Fabrication'", "Order status changes\nSuccess message 'Đã bắt đầu gia công đơn kính'", "Order must be in queue"],
        ["TC_UC38_03", "Complete Fabrication process", "1. Login as Optician\n2. Select an order in fabrication\n3. Click 'Complete Fabrication'", "Order status changes\nSuccess message 'Đã hoàn tất gia công đơn kính'", "Order must be in fabrication state"],
    ],
    "UC-55": [
        ["TC_UC55_01", "Create a new room", "1. Login as Manager\n2. Navigate to Room Management\n3. Click 'Add New Room'\n4. Fill in room details (Name, Category)\n5. Submit", "Room is successfully created", "Must be logged in as Manager"],
        ["TC_UC55_02", "Update an existing room", "1. Login as Manager\n2. Select an existing room\n3. Edit details\n4. Submit", "Room details are updated successfully", "Room must exist"],
        ["TC_UC55_03", "Deactivate a room (soft delete)", "1. Login as Manager\n2. Select an active room\n3. Click 'Deactivate'", "Room status becomes inactive (soft delete)\nRoom is hidden from active room lists", "Room must be active"],
        ["TC_UC55_04", "Reactivate a room", "1. Login as Manager\n2. Select an inactive room\n3. Click 'Reactivate'", "Room status becomes active", "Room must be inactive"],
        ["TC_UC55_05", "View rooms by category", "1. Login as Manager\n2. Navigate to Room Management\n3. Filter by specific category (e.g., EXAM)", "Only rooms belonging to the selected category are displayed", "Rooms of that category must exist"],
    ],
    "UC-56": [
        ["TC_UC56_01", "Assign a staff to a room", "1. Login as Manager\n2. Navigate to Staff Room Roster\n3. Select a date\n4. Select a staff member and a room\n5. Submit assignment", "Staff is successfully assigned to the room for the specified date", "Staff and Room must exist and be active"],
        ["TC_UC56_02", "View Staff Roster for a specific date", "1. Login as Manager\n2. Navigate to Staff Roster\n3. Select a specific date", "List of staff on duty and their assigned rooms for that date is displayed", "Assignments must exist for the date"],
        ["TC_UC56_03", "Resolve Room for Staff (Read-Only)", "1. Login as Receptionist/Doctor\n2. Trigger a check-in or booking flow requiring room resolution\n3. System calls resolve room endpoint", "System successfully fetches the resolved room for the staff without allowing assignment changes", "Staff must have an assigned room"],
    ]
}

file_path = "c:/Users/ADMIN/Clinic-Management-System/ECMS_System_Test.xlsx"
wb = openpyxl.load_workbook(file_path)

headers = ["Test Case ID", "Test Case Description", "Test Case Procedure", "Expected Results", "Pre-Conditions"]
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="4F81BD")
border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

for sheet_name, cases in test_cases.items():
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = 30
        
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 40
        
    for row_num, case in enumerate(cases, 2):
        for col_num, value in enumerate(case, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical="top")

wb.save(file_path)
print("Updated Excel file successfully.")
